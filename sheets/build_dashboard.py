#!/usr/bin/env python3
"""
Build the EDLD mining dashboard template (Mining_Dashboard.xlsx).

The output is meant for Google Sheets: File > Import > Upload, or open it from
Drive. It is not meant to be used in Excel or LibreOffice, which cannot evaluate
the dashboard's formulas.

Why the odd formula wrapper
---------------------------
The dashboard and the dropdown lists use Sheets-only constructs (FILTER, SORT,
UNIQUE, LET, array literals over ranges). An xlsx cell cannot hold those
directly, so they are written the way Sheets writes them when it exports:

    =IFERROR(__xludf.DUMMYFUNCTION("<formula text>"), <cached value>)

Sheets recognises the wrapper on import and restores the formula inside. xlsx
caps a string literal at 255 characters, so the text is split into chunks
joined with &, as Google's own export does.

The formula text itself lives in Dashboard.gs (FORMULAS), which is also what
"Repair dashboard formulas" writes. One definition, two consumers.

Usage:
    python3 sheets/build_dashboard.py [--out PATH] [--version YYYYMMDD]

--version defaults to the repository's version file, so a template built for a
release carries that release's datestamp in Settings!D1.

Requires openpyxl.
"""

import argparse
import datetime as dt
import pathlib
import re
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent

# ── Content ──────────────────────────────────────────────────────────────────

# Positional contract with Code.gs COLUMNS and core/sheets_publish.py.
DEPOSIT_COLUMNS = [
    'deposit_id', 'system', 'system_address', 'body', 'body_id', 'planet_class',
    'gravity', 'body_radius_m', 'atmosphere', 'volcanism', 'signal_no', 'commodity',
    'latitude', 'longitude', 'density_claimed', 'density_observed', 'amount', 'rigs',
    'refine_count', 'first_seen', 'last_confirmed', 'reported_by', 'is_test',
    'depleted_on',
]

# Dashboard table columns, in the order the FILTER in Dashboard.gs emits them.
HEADERS = ['System', 'Body', 'Gravity', 'Mining Site', 'Commodity',
           'Amount', 'Density', 'Rigs', 'Depleted']

PLACEHOLDER_SQUADRON = 'Your Squadron Name [TAG]'
PLACEHOLDER_MAINTAINER = 'Your CMDR Name'

ROLES = [
    ('Background',  'Sheet background'),
    ('Panel',       'Input fields, data rows'),
    ('Panel Alt',   'Alternate data rows'),
    ('Accent',      'HUD stripe, header bar, borders'),
    ('Accent Text', 'Text on the header bar'),
    ('Title',       'Dashboard title'),
    ('Text',        'Data text'),
    ('Text Dim',    "Labels, subtitle, 'Low' values"),
    ('Border',      'Swatch and table rules'),
    ('Highlight',   "'High' amount / density"),
    ('Warning',     'Depleted deposits'),
    ('Input Text',  'Dropdown text'),
]

# Classic HUD must match ED.FALLBACK in Dashboard.gs.
THEMES = {
    'Classic HUD': ['#0A0A0A', '#16100A', '#2B1E0E', '#FF7100', '#0A0A0A', '#FF8C1A',
                    '#FFB266', '#9A5A1E', '#4A2800', '#FFD24A', '#E0301E', '#FFA040'],
    'Federation':  ['#060A12', '#0B1424', '#1A2F52', '#2E8BFF', '#04070D', '#5AA8FF',
                    '#BFD9FF', '#5577A0', '#1C3C66', '#FFFFFF', '#FF4A3D', '#8CC4FF'],
    'Empire':      ['#0C0710', '#170D1D', '#2E1A3A', '#D9A93D', '#140A18', '#F0C75E',
                    '#E9DDC0', '#8C7A9E', '#4A2F59', '#FFE9A8', '#FF5A4F', '#F0C75E'],
    'Alliance':    ['#050B07', '#0A1810', '#173724', '#3DCC6B', '#04100A', '#5EE88A',
                    '#C6F2D3', '#4F8A62', '#1D4D2D', '#EFFFF4', '#FF6A3D', '#7BF09F'],
    'Thargoid':    ['#030A0A', '#071716', '#123634', '#33E0C6', '#031110', '#6FF5DF',
                    '#B8F5EC', '#3E8C82', '#135048', '#D4FF5A', '#FF3D7F', '#6FF5DF'],
}
DEFAULT_THEME = 'Classic HUD'
FONTS = ['Orbitron', 'Exo 2', 'Rajdhani', 'Oxanium', 'Share Tech Mono', 'Arial']
TITLE_FONT, BODY_FONT, MONO_FONT = 'Orbitron', 'Exo 2', 'Share Tech Mono'

SETTINGS_NOTES = [
    'HOW IT WORKS',
    '• Custom Hex (column B) is yours to edit, as #RRGGBB. Its swatch updates when the dashboard script runs.',
    '• Active Hex (column D) is what the dashboard uses: the chosen preset, or column B when Theme = Custom.',
    '• ED Dashboard ▸ Copy active theme to Custom starts you from a preset, so you only change what you want.',
    '• More presets: unhide the Themes tab and add a column before "Custom". It joins the Theme list.',
    '• Colours re-apply on edit once Dashboard.gs is installed (Extensions ▸ Apps Script).',
]

LAST_ROW = 1000
ROLE_ROW0 = 10   # Settings!A10 — ED.ROLE_FIRST_ROW
NOTES_ROW = 24   # ED.NOTES_ROW


# ── Formula plumbing ─────────────────────────────────────────────────────────

def load_formulas(gs_path):
    """Pull FORMULAS out of Dashboard.gs: 'Sheet!A1': `=...`."""
    text = gs_path.read_text(encoding='utf-8')
    block = re.search(r'const FORMULAS = \{(.*?)\n\};', text, re.S)
    if not block:
        sys.exit(f'FORMULAS not found in {gs_path}')
    found = dict(re.findall(r"'([A-Za-z]+![A-Z]+\d+)':\s*`([^`]*)`", block.group(1)))
    for k, v in found.items():
        if '${' in v:
            sys.exit(f'{k}: FORMULAS entries must not interpolate')
    return found


def sheets_only(formula, cached='""', chunk=250):
    """Wrap a Sheets formula the way Sheets exports it, so import restores it."""
    body = formula[1:] if formula.startswith('=') else formula
    parts = [body[i:i + chunk] for i in range(0, len(body), chunk)]
    literal = '&'.join('"' + p.replace('"', '""') + '"' for p in parts)
    return f'=IFERROR(__xludf.DUMMYFUNCTION({literal}),{cached})'


def unwrap(cell_formula):
    """Inverse of sheets_only(), for the self-check."""
    m = re.fullmatch(r'=IFERROR\(__xludf\.DUMMYFUNCTION\((.*)\),.*\)', cell_formula, re.S)
    parts = re.findall(r'"((?:[^"]|"")*)"(?:&|$)', m.group(1))
    return '=' + ''.join(p.replace('""', '"') for p in parts)


# ── Styling helpers ──────────────────────────────────────────────────────────

C = dict(zip([r[0] for r in ROLES], THEMES[DEFAULT_THEME]))


def h(x):
    return x.lstrip('#').upper()


def fill(x):
    return PatternFill('solid', start_color=h(x), end_color=h(x))


def side(x, style='thin'):
    return Side(style=style, color=h(x))


def box(x, bottom='thin'):
    return Border(left=side(x), right=side(x), top=side(x), bottom=side(x, bottom))


def contrast(x):
    r, g, b = (int(h(x)[i:i + 2], 16) for i in (0, 2, 4))
    return '#000000' if (0.299 * r + 0.587 * g + 0.114 * b) > 140 else '#FFFFFF'


def font(color, name=BODY_FONT, **kw):
    return Font(name=name, color=h(color), **kw)


def paint(ws, rows, cols, bg, fg):
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(r, c)
            cell.fill = fill(bg)
            cell.font = font(fg)


# ── Sheets ───────────────────────────────────────────────────────────────────

def build_dashboard(wb, formulas):
    ds = wb.active
    ds.title = 'Dashboard'
    paint(ds, LAST_ROW, 10, C['Background'], C['Text'])

    ds['A1'] = ('=IF(Settings!B1<>"",Settings!B1 & ": Mining Data",'
                'IF(Settings!B2<>"","CMDR " & Settings!B2 & "\'s Mining Data",""))')
    ds['A2'] = '=IF(A1=Settings!B1 & ": Mining Data","Maintained by CMDR: " & Settings!B2,"")'
    ds['E2'] = '=Settings!D1'
    ds.merge_cells('E2:I2')
    ds['A1'].font = font(C['Title'], TITLE_FONT, bold=True, size=20)
    ds['A1'].alignment = Alignment(vertical='center')
    ds['A2'].font = font(C['Text Dim'], italic=True, size=10)
    ds['E2'].font = font(C['Text Dim'], size=9)
    ds['E2'].alignment = Alignment(horizontal='right')
    ds.row_dimensions[1].height = 38
    ds.row_dimensions[3].height = 4
    for c in range(1, 10):
        ds.cell(3, c).fill = fill(C['Accent'])

    # Filters (A:B) and sort (E:F)
    ds['A4'], ds['B4'], ds['E4'], ds['F4'] = 'System', 'Commodity', 'Sort By', 'Order'
    ds['E5'], ds['F5'] = HEADERS[0], 'Ascending'
    for a in ('A4', 'B4', 'E4', 'F4'):
        ds[a].font = font(C['Text Dim'], bold=True, size=9)
    for a in ('A5', 'B5', 'E5', 'F5'):
        ds[a].fill = fill(C['Panel'])
        ds[a].font = font(C['Input Text'], bold=True)
        ds[a].border = box(C['Accent'], 'medium')
        ds[a].alignment = Alignment(indent=1, vertical='center')
    ds.row_dimensions[5].height = 22
    ds.row_dimensions[7].height = 24

    for sqref, f1 in (('A5', 'Queries!$A$1:$A$1000'), ('B5', 'Queries!$B$1:$B$1000'),
                      ('E5', 'Queries!$D$1:$L$1'), ('F5', '"Ascending,Descending"')):
        dv = DataValidation(type='list', formula1=f1, allow_blank=True)
        ds.add_data_validation(dv)
        dv.add(sqref)

    # The table: one formula at A7 fills A7:I
    ds['A7'] = sheets_only(formulas['Dashboard!A7'])

    left = (1, 2, 5)
    for r in range(7, LAST_ROW + 1):
        for c in range(1, 10):
            ds.cell(r, c).alignment = Alignment(
                horizontal='left' if c in left else 'center',
                vertical='center' if r == 7 else None,
                indent=1 if c in left else 0)
        if r > 7:
            ds.cell(r, 3).number_format = '0.000" g"'
            ds.cell(r, 4).number_format = '0'
            ds.cell(r, 8).number_format = '0'

    # Conditional formatting. Sheets imports xlsx rules grouped by range, in the
    # order each range first appears, and applies the first rule that matches —
    # so ranges are added most-specific first, and rules that share a range are
    # written to be mutually exclusive. Dashboard.gs rebuilds the same set.
    def rule(formula, bg, fg, **kw):
        return FormulaRule(formula=[formula], stopIfTrue=True,
                           font=Font(color=h(fg), **kw),
                           fill=PatternFill('solid', start_color=h(bg),
                                            end_color=h(bg), bgColor=h(bg)))

    cf = ds.conditional_formatting
    cf.add('A7:I7', rule('$A$7<>""', C['Accent'], C['Accent Text'], bold=True))
    bands = ((0, C['Panel']), (1, C['Panel Alt']))   # row 8, the first data row, is even
    hilo, allr = f'F8:G{LAST_ROW}', f'A8:I{LAST_ROW}'
    for p, bg in bands:
        cf.add(hilo, rule(f'AND($A8<>"",$I8="",F8="High",MOD(ROW(),2)={p})', bg, C['Highlight'], bold=True))
        cf.add(hilo, rule(f'AND($A8<>"",$I8="",F8="Low",MOD(ROW(),2)={p})', bg, C['Text Dim']))
    for p, bg in bands:
        cf.add(allr, rule(f'AND($A8<>"",$I8<>"",MOD(ROW(),2)={p})', bg, C['Warning'],
                          italic=True, strike=True))
    for p, bg in bands:
        cf.add(allr, rule(f'AND($A8<>"",$I8="",MOD(ROW(),2)={p})', bg, C['Text']))

    for col, w in zip('ABCDEFGHIJ', [18, 22, 11, 14, 24, 13, 11, 8, 14, 3]):
        ds.column_dimensions[col].width = w
    ds.column_dimensions.group('K', 'Z', hidden=True)
    ds.sheet_view.showGridLines = False
    ds.freeze_panes = 'A8'
    ds.sheet_properties.tabColor = h(C['Accent'])
    return ds


def build_settings(wb, version):
    st = wb.create_sheet('Settings')
    paint(st, 40, 8, C['Background'], C['Text'])
    st['A1'], st['B1'] = 'Squadron', PLACEHOLDER_SQUADRON
    st['A2'], st['B2'] = 'Maintainer', PLACEHOLDER_MAINTAINER
    st['D1'] = f'Mining Dashboard v{version} powered by EDLD'
    st.merge_cells('D1:H1')
    st['D1'].font = font(C['Text Dim'], size=9)
    st['D1'].alignment = Alignment(horizontal='right')

    st['A4'] = 'DASHBOARD THEME'
    st['A5'], st['B5'] = 'Theme', DEFAULT_THEME
    st['A6'], st['B6'] = 'Title Font', TITLE_FONT
    st['A7'], st['B7'] = 'Body Font', BODY_FONT
    st['C5'] = 'Pick a preset, or "Custom" to use your own hex codes below'
    st['C6'] = 'Google Fonts — any Sheets font name works'

    for j, t in enumerate(['Color Role', 'Custom Hex', '', 'Active Hex', '', 'Used For']):
        cell = st.cell(ROLE_ROW0 - 1, 1 + j, t)
        cell.fill = fill(C['Accent'])
        cell.font = font(C['Accent Text'], bold=True)
    for i, (role, desc) in enumerate(ROLES):
        r = ROLE_ROW0 + i
        hexv = THEMES[DEFAULT_THEME][i]
        band = C['Panel'] if i % 2 == 0 else C['Panel Alt']
        st.cell(r, 1, role)
        st.cell(r, 2, hexv)
        st.cell(r, 4, f'=IFERROR(INDEX(Themes!$B$2:$Z$50,MATCH($A{r},Themes!$A$2:$A$50,0),'
                      f'MATCH($B$5,Themes!$B$1:$Z$1,0)),B{r})')
        st.cell(r, 6, desc)
        for c in (1, 2, 4, 6):
            st.cell(r, c).fill = fill(band)
            st.cell(r, c).border = Border(bottom=side(C['Border']))
        for c in (3, 5):
            st.cell(r, c).fill = fill(hexv)
            st.cell(r, c).border = box(C['Border'])
        st.cell(r, 2).font = font(C['Input Text'], MONO_FONT, bold=True)
        st.cell(r, 4).font = font(C['Text'], MONO_FONT)
        st.cell(r, 6).font = font(C['Text Dim'], italic=True)
        st.cell(r, 2).alignment = st.cell(r, 4).alignment = Alignment(horizontal='center')

    for k, t in enumerate(SETTINGS_NOTES):
        st.cell(NOTES_ROW + k, 1, t).font = font(C['Text Dim'], size=10)
    st.cell(NOTES_ROW, 1).font = font(C['Title'], TITLE_FONT, bold=True, size=10)

    for a in ('A1', 'A2', 'A5', 'A6', 'A7'):
        st[a].font = font(C['Text Dim'], bold=True)
    for a in ('B1', 'B2', 'B5', 'B6', 'B7'):
        st[a].fill = fill(C['Panel'])
        st[a].font = font(C['Input Text'], bold=True)
        st[a].border = Border(bottom=side(C['Accent']))
    for a in ('C5', 'C6'):
        st[a].font = font(C['Text Dim'], italic=True, size=9)
    st['A4'].font = font(C['Title'], TITLE_FONT, bold=True, size=12)

    dv_theme = DataValidation(type='list', formula1='Themes!$B$1:$Z$1', allow_blank=False)
    dv_font = DataValidation(type='list', formula1='"' + ','.join(FONTS) + '"',
                             allow_blank=False, showErrorMessage=False)
    dv_hex = DataValidation(
        type='custom',
        formula1=f'AND(LEN(B{ROLE_ROW0})=7,LEFT(B{ROLE_ROW0},1)="#",'
                 f'ISNUMBER(HEX2DEC(MID(B{ROLE_ROW0},2,6))))',
        showErrorMessage=True, errorTitle='Hex colour',
        error='Enter a colour as #RRGGBB, e.g. #FF7100')
    for dv, ref in ((dv_theme, 'B5'), (dv_font, 'B6:B7'),
                    (dv_hex, f'B{ROLE_ROW0}:B{ROLE_ROW0 + len(ROLES) - 1}')):
        st.add_data_validation(dv)
        dv.add(ref)

    for col, w in zip('ABCDEFGH', [18, 36, 5, 12, 5, 34, 4, 4]):
        st.column_dimensions[col].width = w
    st.sheet_view.showGridLines = False
    st.sheet_properties.tabColor = h(C['Text Dim'])


def build_deposits(wb):
    dp = wb.create_sheet('Deposits')
    for c, name in enumerate(DEPOSIT_COLUMNS, 1):
        dp.cell(1, c, name).font = Font(name='Arial', bold=True)
    lat = DEPOSIT_COLUMNS.index('latitude') + 1
    for r in range(2, LAST_ROW + 1):
        for c in (lat, lat + 1):
            dp.cell(r, c).number_format = '0.000000'
    dp.freeze_panes = 'A2'
    dp.sheet_state = 'hidden'


def build_queries(wb, formulas):
    q = wb.create_sheet('Queries')
    q['A1'] = sheets_only(formulas['Queries!A1'])
    q['B1'] = sheets_only(formulas['Queries!B1'])
    # D1:L1 header names (the Sort By list); D2:L2 the same with a sort arrow.
    for j, name in enumerate(HEADERS):
        col = get_column_letter(4 + j)
        q[f'{col}1'] = name
        q[f'{col}2'] = (f'={col}1&IF({col}1=Dashboard!$E$5,'
                        f'IF(Dashboard!$F$5="Descending"," ▼"," ▲"),"")')
    q['N1'] = '=IFERROR(MATCH(Dashboard!$E$5,$D$1:$L$1,0),1)'
    q['N2'] = '=Dashboard!$F$5<>"Descending"'
    q['O1'] = '← sort column index'
    q['O2'] = '← ascending?'
    q.sheet_state = 'hidden'


def build_themes(wb):
    th = wb.create_sheet('Themes')
    names = list(THEMES) + ['Custom']
    th['A1'] = 'Color Role'
    for j, n in enumerate(names):
        th.cell(1, 2 + j, n)
    for cell in th[1]:
        cell.fill = fill(C['Accent'])
        cell.font = font(C['Accent Text'], bold=True)
        cell.alignment = Alignment(horizontal='center')
    for i, (role, _) in enumerate(ROLES):
        r = 2 + i
        th.cell(r, 1, role).font = font(C['Text'])
        th.cell(r, 1).fill = fill(C['Panel'])
        for j, n in enumerate(THEMES):
            v = THEMES[n][i]
            cell = th.cell(r, 2 + j, v)
            cell.fill = fill(v)
            cell.font = font(contrast(v))
            cell.alignment = Alignment(horizontal='center')
        cc = th.cell(r, 2 + len(THEMES), f'=Settings!$B${ROLE_ROW0 + i}')
        cc.fill = fill(C['Panel'])
        cc.font = font(C['Text'])
        cc.alignment = Alignment(horizontal='center')
    note = th.cell(len(ROLES) + 3, 1,
                   'Add your own scheme: insert a column before "Custom", name it in row 1 and put '
                   'hex codes below. It joins the Settings theme list. "Custom" mirrors '
                   f'Settings!B{ROLE_ROW0}:B{ROLE_ROW0 + len(ROLES) - 1}; edit it there.')
    note.font = font(C['Text Dim'], italic=True)
    th.column_dimensions['A'].width = 16
    for j in range(len(names)):
        th.column_dimensions[get_column_letter(2 + j)].width = 14
    th.sheet_view.showGridLines = False
    th.freeze_panes = 'B2'
    th.sheet_properties.tabColor = h(C['Text Dim'])
    th.sheet_state = 'hidden'


def default_version():
    vf = ROOT / 'version'
    if vf.is_file():
        v = vf.read_text(encoding='utf-8').strip()
        if re.fullmatch(r'\d{8}', v):
            return v
    return dt.date.today().strftime('%Y%m%d')


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('--out', default=str(HERE / 'Mining_Dashboard.xlsx'))
    ap.add_argument('--version', default=default_version())
    ap.add_argument('--gs', default=str(HERE / 'Dashboard.gs'))
    args = ap.parse_args()

    formulas = load_formulas(pathlib.Path(args.gs))
    for key in ('Dashboard!A7', 'Queries!A1', 'Queries!B1'):
        if key not in formulas:
            sys.exit(f'{key} missing from FORMULAS in {args.gs}')

    wb = Workbook()
    build_dashboard(wb, formulas)
    build_settings(wb, args.version)
    build_deposits(wb)
    build_queries(wb, formulas)
    build_themes(wb)

    # Self-check: every wrapped formula must unwrap to exactly what Dashboard.gs says.
    for key, src in formulas.items():
        sheet, a1 = key.split('!')
        if unwrap(wb[sheet][a1].value) != src:
            sys.exit(f'{key}: wrapped formula does not round-trip')

    wb.active = 0
    wb.save(args.out)
    print(f'wrote {args.out}')


if __name__ == '__main__':
    main()
